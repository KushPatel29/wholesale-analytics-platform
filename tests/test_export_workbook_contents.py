"""
What a downloaded workbook actually shows, read back out of the file.

The suite had 118 tests over the export paths and none of them opened the
workbook, so a margin of 22.96% had been downloading as **2,296%** in every
XLSX the app produced. Excel's own `0.0%` format multiplies the stored number
by 100 before displaying it, and this app stores percentages on the 0-100
scale (`app/services/formatting.py`: "a percentage that is already on the 0-100
scale"). Asserting on the DataFrame that goes *into* the writer cannot catch
that - the number is right and the file is wrong - so every test here reads the
cell back and reconstructs what Excel will paint.
"""

from __future__ import annotations

import html
import io
import re
import zipfile

import openpyxl
import pandas as pd
import pytest

from app.core.exports import (
    CURRENCY_FORMAT,
    PERCENT_FORMAT,
    column_unit,
    dataframes_to_xlsx_bytes,
)


@pytest.fixture(autouse=True)
def permitted_viewer(monkeypatch):
    """Export as someone allowed to see cost and margin.

    Without this the RBAC layer blanks every cost and margin column before the
    writer sees it, and a test asserting on the format of an all-`None` column
    would pass whatever the formatter did. `test_masking_still_blanks_costs`
    covers the other side.
    """
    monkeypatch.setattr(
        "app.core.sensitive_data.sensitive_access_flags",
        lambda user=None: {
            "cost": True,
            "margin": True,
            "profit": True,
            "recommendations": True,
            "margin_risk": True,
            "export_sensitive": True,
        },
    )


def _sheet(frame: pd.DataFrame, name: str = "Data"):
    data = dataframes_to_xlsx_bytes({name: frame})
    return openpyxl.load_workbook(io.BytesIO(data))[name], data


def _formats(frame: pd.DataFrame, name: str = "Data") -> dict[str, str]:
    """Build the sheet and return header -> number format."""
    data = dataframes_to_xlsx_bytes({name: frame})
    return _column_formats(data, list(frame.columns))


def _column_formats(payload: bytes, columns) -> dict[str, str]:
    """Map header -> the number format Excel will apply to that column.

    Read out of the sheet XML rather than through
    `worksheet.column_dimensions`. xlsxwriter collapses adjacent columns that
    share a width and a format into a single `<col min="1" max="2">` span, and
    openpyxl registers such a span only under its first letter - so every
    second column of a pair reads as "General" whatever was written to it. An
    earlier version of this helper had that flaw and called four correctly
    formatted percentage columns unformatted.
    """
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        sheet_xml = archive.read("xl/worksheets/sheet1.xml").decode("utf-8")
        styles_xml = archive.read("xl/styles.xml").decode("utf-8")

    custom = {
        fmt_id: html.unescape(code)
        for fmt_id, code in re.findall(r'<numFmt numFmtId="(\d+)" formatCode="([^"]*)"/>', styles_xml)
    }
    builtin = {"0": "General", "3": "#,##0", "4": "#,##0.00"}
    cell_xfs = re.search(r"<cellXfs[^>]*>(.*?)</cellXfs>", styles_xml, re.S)
    xfs = re.findall(r'<xf [^>]*numFmtId="(\d+)"', cell_xfs.group(1) if cell_xfs else "")

    by_index: dict[int, str] = {}
    for tag in re.findall(r"<col [^>]*/>", sheet_xml):
        low = int(re.search(r'min="(\d+)"', tag).group(1))
        high = int(re.search(r'max="(\d+)"', tag).group(1))
        style_match = re.search(r'style="(\d+)"', tag)
        code = "General"
        if style_match and int(style_match.group(1)) < len(xfs):
            fmt_id = xfs[int(style_match.group(1))]
            code = custom.get(fmt_id) or builtin.get(fmt_id, f"numFmtId:{fmt_id}")
        for index in range(low, high + 1):
            by_index[index] = code

    return {str(name): by_index.get(index, "General") for index, name in enumerate(columns, start=1)}


def _displayed(value: float, number_format: str) -> str:
    """The string Excel paints for a stored value under a number format.

    Excel's rule is about the *unquoted* percent sign: a bare `%` in a format
    scales the cell by 100, while a `%` inside quotes is just a character to
    print. That distinction is the entire bug, so it is derived from the format
    string here rather than compared against this module's own constant - a
    comparison against the constant passes no matter which format ships.
    """
    unquoted = re.sub(r'"[^"]*"', "", number_format)
    if "%" in unquoted:
        return f"{value * 100:.1f}%"
    if "%" in number_format:
        return f"{value:.1f}%"
    if number_format == CURRENCY_FORMAT:
        return f"${value:,.2f}"
    return str(value)


class TestPercentagesMatchTheDashboard:
    def test_margin_of_22_96_displays_as_23_0_percent(self):
        """The reported bug: 22.96 must not paint as 2,296%."""
        frame = pd.DataFrame({"margin_pct": [22.96]})
        worksheet, _ = _sheet(frame)
        stored = worksheet.cell(row=2, column=1).value
        number_format = _formats(frame)["margin_pct"]

        assert stored == pytest.approx(22.96), "the stored number must still be the page's number"
        assert _displayed(stored, number_format) == "23.0%"

    def test_percent_format_does_not_multiply(self):
        assert "%" in PERCENT_FORMAT
        assert PERCENT_FORMAT != "0.0%", "the built-in percent format multiplies the cell by 100"

    @pytest.mark.parametrize(
        "column",
        [
            "margin_pct",
            "revenue_share_pct",
            "cost_coverage_pct",
            "yoy_revenue_pct",
            "otif_pct",
            "Cumulative%",
            "Premium Share %",
            "MarginPct",
        ],
    )
    def test_percent_columns_are_percent_formatted(self, column):
        """`revenue_share_pct` used to print as `$23.40`.

        The currency test ran before the percent test, so any percentage whose
        name mentioned money was formatted as money.
        """
        assert _formats(pd.DataFrame({column: [23.4]}))[column] == PERCENT_FORMAT
        assert _displayed(23.4, PERCENT_FORMAT) == "23.4%"


class TestUnitsThatOnlyLookLikePercentages:
    @pytest.mark.parametrize("column", ["blended_rate", "Blended Rate", "Effective Rate"])
    def test_pay_rates_are_currency(self, column):
        """`blended_rate` is labor cost over paid hours - dollars per hour."""
        assert _formats(pd.DataFrame({column: [27.55]}))[column] == CURRENCY_FORMAT

    @pytest.mark.parametrize("column", ["generated_at", "generated_at_utc"])
    def test_timestamps_are_not_percentages(self, column):
        """"gene*rate*d_at" matched the old substring test for "rate"."""
        assert column_unit(column) != "percent"

    def test_revenue_is_currency(self):
        assert _formats(pd.DataFrame({"Revenue": [125000.0]}))["Revenue"] == CURRENCY_FORMAT


class TestMaskedColumnsDoNotBreakTheSheet:
    def test_masking_still_blanks_costs_for_a_viewer_without_permission(self, monkeypatch):
        """The other side of `permitted_viewer` - masking is still enforced."""
        monkeypatch.setattr(
            "app.core.sensitive_data.sensitive_access_flags",
            lambda user=None: {
                "cost": False,
                "margin": False,
                "profit": False,
                "recommendations": False,
                "margin_risk": False,
                "export_sensitive": False,
            },
        )
        worksheet, _ = _sheet(pd.DataFrame({"Revenue": [125000.0], "margin_pct": [22.96]}))

        assert worksheet.cell(row=2, column=1).value == 125000.0
        assert worksheet.cell(row=2, column=2).value is None

    def test_a_masked_column_does_not_strip_formats_from_later_columns(self):
        """Permission masking blanks a whole column.

        Column widths and formats were applied inside one `try`, so the first
        column that raised - an all-`None` column makes the width heuristic
        raise on `int(NaN)` - silently left every column after it unformatted.
        """
        frame = pd.DataFrame(
            {
                "Revenue": [125000.0],
                "unit_cost": [None],
                "margin_pct": [22.96],
                "OrderCount": [412],
            }
        )
        formats = _formats(frame)

        assert formats["margin_pct"] == PERCENT_FORMAT
        assert formats["Revenue"] == CURRENCY_FORMAT


class TestExportedTextStaysText:
    @pytest.mark.parametrize(
        "text",
        ["=1+1", '=HYPERLINK("http://example.test","click")', "=Rebate Program", "+1 Vendor"],
    )
    def test_leading_equals_is_not_a_formula(self, text):
        """A product literally called "=Rebate Program" is a name, not a formula."""
        worksheet, payload = _sheet(pd.DataFrame({"Product": [text]}))
        cell = worksheet.cell(row=2, column=1)

        assert cell.data_type == "s", f"{text!r} was written as a formula cell"
        assert cell.value == text, "the exported text must read back unchanged"

        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            sheet_xml = archive.read("xl/worksheets/sheet1.xml").decode("utf-8")
        assert "<f>" not in sheet_xml


class TestColumnUnitClassification:
    @pytest.mark.parametrize(
        ("name", "unit"),
        [
            ("margin_pct", "percent"),
            ("revenue_share_pct", "percent"),
            ("cost_coverage_pct", "percent"),
            ("repeat_rate", "percent"),
            ("Rate Delta %", "percent"),
            # Naming what a proportion is *of* does not make it money.
            ("core_revenue_share", "percent"),
            ("below_target_revenue_share", "percent"),
            ("cost_null_rate", "percent"),
            ("effective_tax_rate", "percent"),
            # Rates that really are dollars per unit.
            ("blended_rate", "currency"),
            ("department_rate", "currency"),
            ("Revenue", "currency"),
            ("total_cost", "currency"),
            ("AvgPrice", "currency"),
            # Neither: `current_ratio` is 1.8, not 1.8%.
            ("current_ratio", "count"),
            ("cancelled_ratio", "count"),
            ("OrderCount", "count"),
            ("generated_at", "count"),
        ],
    )
    def test_unit(self, name, unit):
        assert column_unit(name) == unit


class TestLaborPercentagesReachTheWorkbookOnTheHouseScale:
    """Labor stores proportions as fractions; the workbook writer assumes 0-100.

    Every `*_pct` in `labor_bundle` is a plain ratio (`premium_cost /
    labor_cost`) and the page multiplies on the way to the screen. Left alone,
    a page reading "3.2%" downloaded as "0.0%", so the frames are put on the
    house scale at labor's export boundary.
    """

    def test_fraction_columns_are_scaled_to_points(self):
        from app.services.labor_bundle import _li_percent_columns_to_points

        frames = {
            "Departments": pd.DataFrame(
                {
                    "Department": ["Grocery"],
                    "Labor Cost": [125000.0],
                    "Blended Rate": [27.55],
                    "Premium Share %": [0.1234],
                    "Rate Delta %": [0.032],
                    "Cost Volatility": [0.41],
                }
            )
        }
        out = _li_percent_columns_to_points(frames)["Departments"]

        assert out["Premium Share %"].iloc[0] == pytest.approx(12.34)
        assert out["Rate Delta %"].iloc[0] == pytest.approx(3.2)
        # A pay rate is dollars per hour, and a volatility is not a percentage.
        assert out["Blended Rate"].iloc[0] == pytest.approx(27.55)
        assert out["Cost Volatility"].iloc[0] == pytest.approx(0.41)
        assert out["Labor Cost"].iloc[0] == pytest.approx(125000.0)

    def test_summary_sheet_scales_by_its_format_column(self):
        """The Summary sheet is one metric column; the unit is a sibling cell."""
        from app.services.labor_bundle import _li_percent_columns_to_points

        frames = {
            "Summary": pd.DataFrame(
                {
                    "Metric": ["Premium Share", "Total Labor Cost", "Window Start"],
                    "Value": [0.1234, 125000.0, "2026-01-05"],
                    "Format": ["percent", "currency", "date"],
                }
            )
        }
        out = _li_percent_columns_to_points(frames)["Summary"]

        assert float(out["Value"].iloc[0]) == pytest.approx(12.34)
        assert float(out["Value"].iloc[1]) == pytest.approx(125000.0)
        assert out["Value"].iloc[2] == "2026-01-05"

    def test_empty_frames_pass_through(self):
        from app.services.labor_bundle import _li_percent_columns_to_points

        out = _li_percent_columns_to_points({"Empty": pd.DataFrame()})
        assert out["Empty"].empty


class TestUnitsThatAreNeitherMoneyNorPercentage:
    """A dispersion measure carries no unit, and the words around it lie.

    `Rate Volatility` was printed as a percentage on the strength of "rate"
    and `Cost Volatility` as dollars on the strength of "cost", so the same
    statistic appeared twice in one sheet in two different wrong units.
    """

    @pytest.mark.parametrize(
        "column",
        ["Rate Volatility", "Cost Volatility", "cost_volatility", "customer_hhi", "priority_score"],
    )
    def test_dispersion_is_neither(self, column):
        assert column_unit(column) == "count"

    def test_a_fraction_keeps_its_decimals(self):
        """`#,##0` would print a coefficient of variation of 0.41 as "0"."""
        formats = _formats(pd.DataFrame({"Cost Volatility": [0.41, 0.38]}))
        assert formats["Cost Volatility"] == "#,##0.00"

    def test_whole_numbers_stay_whole(self):
        formats = _formats(pd.DataFrame({"OrderCount": [412, 87]}))
        assert formats["OrderCount"] == "#,##0"


class TestObjectDtypeColumnsStillGetTheirUnit:
    def test_a_percentage_carrying_nulls_is_still_formatted(self):
        """A frame holding `pd.NA` beside floats is object dtype.

        `is_numeric_dtype` is False for those, so several percentage columns
        shipped unformatted and printed as `3.106348459681461`. What matters is
        whether the values are numbers, not what dtype pandas settled on.
        """
        frame = pd.DataFrame({"Absence Share %": [3.106, None, 4.2]}).astype(object)
        assert _formats(frame)["Absence Share %"] == PERCENT_FORMAT

    def test_an_all_null_column_gets_no_unit(self):
        """A masked column is empty; a unit on it would imply a value."""
        frame = pd.DataFrame({"Revenue": [1.0], "unit_cost": [None]})
        assert _formats(frame)["unit_cost"] == "General"
