"""Read back every sheet and row after XLSX name normalization and streaming."""
from __future__ import annotations

import io
import re
from types import SimpleNamespace

import openpyxl
import pandas as pd
import pytest
from flask import Flask

from app.core import exports

NAME_CASES = [
    ["Sales [Q1]*?\\", "'quoted'", ""],
    ["Summary", "summary"],
    ["A" * 31 + "-one", "A" * 31 + "-two"],
    ["Sales/A", "Sales:A", "Sales_A"],
    ["B" * 31, "B" * 31 + "-two", "B" * 28 + "_2"],
]


@pytest.fixture(autouse=True)
def export_viewer(monkeypatch):
    monkeypatch.setattr(exports, "current_user", None)
    monkeypatch.setattr("app.core.sensitive_data.sensitive_access_flags", lambda user=None: {
        "cost": True, "margin": True, "profit": True, "recommendations": True,
        "margin_risk": True, "export_sensitive": True,
    })


def frames_for(names):
    return {name: pd.DataFrame({"marker": [f"record-{index}"], "units": [index + 1]})
            for index, name in enumerate(names)}


def assert_all_records(payload, expected):
    book = openpyxl.load_workbook(io.BytesIO(payload))
    try:
        assert len(book.sheetnames) == len(expected)
        assert len({name.casefold() for name in book.sheetnames}) == len(expected)
        for name in book.sheetnames:
            assert 0 < len(name) <= 31 and not re.search(r"[\[\]:*?/\\]", name)
            assert not name.startswith("'") and not name.endswith("'")
        assert [sheet['A2'].value for sheet in book] == [f"record-{i}" for i in range(len(expected))]
    finally:
        book.close()


@pytest.mark.parametrize("engine", ["xlsxwriter", "openpyxl", None])
@pytest.mark.parametrize("names", NAME_CASES, ids=["punctuation", "case", "length", "sanitized", "suffix"])
def test_in_memory_workbook_preserves_all_sheets_and_rows(monkeypatch, engine, names):
    monkeypatch.setattr(exports, "_xlsx_engine_name", lambda: engine)
    assert_all_records(exports.dataframes_to_xlsx_bytes(frames_for(names)), names)


@pytest.mark.parametrize("engine", ["xlsxwriter", "openpyxl"])
def test_streamed_workbook_preserves_colliding_and_invalid_names(monkeypatch, engine):
    monkeypatch.setattr(exports, "_xlsx_engine_name", lambda: engine)
    names = ["Q" * 31 + "-one", "Q" * 31 + "-two", "Sales [Q1]", "Summary", "summary"]
    app = Flask(__name__)
    app.config['TESTING'] = True

    @app.get('/download')
    def download():
        return exports.dataframes_to_xlsx_response(frames_for(names), threshold_rows=0)

    with app.test_client() as client:
        response = client.get('/download', buffered=True)
        try:
            assert response.status_code == 200
            assert_all_records(response.data, names)
        finally:
            response.close()


@pytest.mark.parametrize("mode", ["csv", "xlsxwriter", "openpyxl"])
def test_temporary_download_exists_until_stream_close_then_is_removed(monkeypatch, tmp_path, mode):
    real_temporary_file = exports.tempfile.NamedTemporaryFile
    monkeypatch.setattr(exports.tempfile, "NamedTemporaryFile",
                        lambda *args, **kwargs: real_temporary_file(*args, dir=tmp_path, **kwargs))
    monkeypatch.setattr(exports, "_xlsx_engine_name", lambda: mode if mode != 'csv' else 'xlsxwriter')
    app = Flask(__name__)
    app.config['TESTING'] = True

    @app.get('/download')
    def download():
        frame = frames_for(['Data'])['Data']
        if mode == 'csv':
            return exports.dataframe_to_csv_response(frame)
        return exports.dataframes_to_xlsx_response({'Data': frame}, threshold_rows=0)

    with app.test_client() as client:
        response = client.get('/download', buffered=False)
        try:
            assert response.status_code == 200
            assert len(list(tmp_path.iterdir())) == 1
            payload = response.get_data()
            assert payload
        finally:
            response.close()
    assert not list(tmp_path.iterdir())


@pytest.mark.parametrize("name", ["Instructions", "instructions"])
def test_instructions_never_overwrite_data_sheet(name):
    book = openpyxl.load_workbook(io.BytesIO(exports.to_excel_bytes(
        pd.DataFrame({"marker": ["record-0"]}), sheet_name=name, instructions=['Read me'],
    )))
    try:
        assert len(book.sheetnames) == 2
        assert book.worksheets[0]['A2'].value == 'record-0'
        assert book.worksheets[1]['A1'].value == 'Read me'
    finally:
        book.close()


def test_chart_uses_the_retained_data_sheet_after_name_sanitizing(monkeypatch):
    monkeypatch.setattr(exports, '_xlsx_engine_name', lambda: 'xlsxwriter')
    raw_name = 'Sales/Q1[' + 'long-name-' * 4
    sheets = {raw_name: pd.DataFrame({'label': ['A'], 'units': [7]})}
    payload = exports.dataframes_to_xlsx_bytes(sheets, chart_specs=[{
        'source_sheet': raw_name, 'chart_sheet': 'Charts[Audit]',
        'category_col': 'label', 'value_cols': ['units'], 'chart_type': 'column',
    }])
    book = openpyxl.load_workbook(io.BytesIO(payload))
    try:
        assert len(book.worksheets) == 2
        assert book.worksheets[1]._charts
        reference = book.worksheets[1]._charts[0].series[0].val.numRef.f
        assert book.worksheets[0].title in reference and '$B$2' in reference
    finally:
        book.close()


def test_assistant_column_policy_preserves_colliding_input_sheets(monkeypatch):
    from app.assistant import tools

    monkeypatch.setattr(tools, '_apply_export_column_policy',
                        lambda ctx, frame, **kwargs: (frame, {'selected_columns': list(frame.columns)}))
    names = ['A' * 31 + '-one', 'A' * 31 + '-two']
    sheets, policies = tools._sheets_with_column_policy(SimpleNamespace(), frames_for(names))
    assert len(sheets) == len(policies) == 2
    assert [frame['marker'].iloc[0] for frame in sheets.values()] == ['record-0', 'record-1']
